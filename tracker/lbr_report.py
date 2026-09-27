"""Local Business Radar: the report's data, and its exports.

build() turns a finished run and its businesses into the one JSON object the
report page draws from: only the fields the page shows, so a 500-business
run stays a few megabytes rather than carrying every raw finding. csv_text()
and xlsx_bytes() export the same leads for a CRM or a sales sheet.
"""

import csv
import io

from tracker import lbr_score

SERVICE_LABELS = lbr_score.SERVICES
WEB_LABELS = {"none": "No website", "social_only": "Social page only", "listing_only": "Directory page only",
              "google_site_retired": "Retired business.site", "dead": "Site down", "broken": "Site broken",
              "parked": "Domain parked", "ssl_error": "Broken certificate", "ok": "Working site",
              "blocked": "Not checked"}


def _slim_business(pid, rank, d):
    prof = d.get("profile") or {}
    disc = d.get("discovery") or {}
    out = {"id": pid, "rank": rank, "name": prof.get("name") or "", "category": prof.get("category") or "",
           "address": prof.get("address") or "", "lat": prof.get("lat"), "lng": prof.get("lng"),
           "phone": prof.get("phone") or "", "website": prof.get("website") or "",
           "maps_url": prof.get("maps_url") or "", "rating": prof.get("rating"), "reviews": prof.get("reviews") or 0,
           "locality": prof.get("locality") or "", "researched": bool(disc.get("selected")),
           "chain": bool(disc.get("chain")), "chain_reason": disc.get("chain_reason") or "",
           "locations": prof.get("locations") or 1, "service_area_only": bool(prof.get("service_area_only"))}
    if not out["researched"]:
        return out
    sc = d.get("score") or {}
    out["score"] = {"total": sc.get("total"), "tier": sc.get("tier"), "need": sc.get("need") or {},
                    "ability": sc.get("ability"), "top": sc.get("top_service"),
                    "ranked": sc.get("services_ranked") or []}
    g = d.get("gbp") or {}
    out["gbp"] = {"score": g.get("score"), "claimed": g.get("claimed"), "photos": g.get("photos"),
                  "checks": [{k: c.get(k) for k in ("key", "label", "status", "detail")} for c in g.get("checks") or []],
                  "unchecked": g.get("unchecked") or []}
    w = d.get("website") or {}
    out["web"] = {"kind": w.get("kind"), "label": WEB_LABELS.get(w.get("kind"), w.get("kind") or ""),
                  "score": w.get("score"), "issues": w.get("issues") or [], "checks": w.get("checks") or {},
                  "tags": w.get("tags") or [], "emails": w.get("emails") or [], "socials": w.get("socials") or {},
                  "builder": w.get("builder") or "", "speed": w.get("speed"), "runs_ads": w.get("runs_ads"),
                  "url": w.get("final_url") or w.get("link") or "", "needs_site": bool(w.get("needs_site")),
                  "copyright_year": w.get("copyright_year")}
    r = d.get("reviews") or {}
    out["rev"] = {"stats": r.get("stats") or {}, "themes": r.get("themes"), "recent": r.get("recent") or [],
                  "topics": r.get("topics") or [], "error": r.get("error")}
    v = d.get("visibility") or {}
    out["vis"] = {"rank": v.get("rank"), "in_pack": v.get("in_pack"), "note": v.get("rank_note") or "",
                  "competitors": v.get("competitors") or [], "ads": v.get("ads"), "market": v.get("market") or {}}
    out["facts"] = d.get("facts") or []
    out["pitch"] = d.get("pitch") or {}
    return out


def build(status, rows):
    """The report payload: the run's headline status plus every business."""
    businesses = [_slim_business(r["place_id"], r.get("rank"), r.get("data") or {}) for r in rows]
    businesses.sort(key=lambda b: (not b["researched"], b["rank"] if b["rank"] is not None else 10 ** 6,
                                   -(b["reviews"] or 0)))
    plan = status.get("plan") or {}
    return {"run": {k: status.get(k) for k in ("id", "status", "label", "area", "focus", "cap", "cost", "summary",
                                               "created_at", "finished_at", "purged", "counts", "error")},
            "area": plan.get("area") or {}, "business": plan.get("business") or {},
            "focus_label": plan.get("focus_label") or "", "estimate": plan.get("estimate") or {},
            "services": SERVICE_LABELS, "web_labels": WEB_LABELS, "businesses": businesses}


# ── Exports ──────────────────────────────────────────────────────────────────
COLUMNS = [
    ("Rank", lambda b: b.get("rank") or ""),
    ("Tier", lambda b: (b.get("score") or {}).get("tier") or ""),
    ("Opportunity score", lambda b: (b.get("score") or {}).get("total")),
    ("Business", lambda b: b["name"]),
    ("Category", lambda b: b["category"]),
    ("Address", lambda b: b["address"]),
    ("Phone", lambda b: b["phone"]),
    ("Emails", lambda b: ", ".join((b.get("web") or {}).get("emails") or [])),
    ("Website", lambda b: b["website"]),
    ("Website finding", lambda b: (b.get("web") or {}).get("label") or ""),
    ("Google rating", lambda b: b.get("rating")),
    ("Reviews", lambda b: b.get("reviews")),
    ("Unanswered 1-2 star reviews", lambda b: ((b.get("rev") or {}).get("stats") or {}).get("unanswered_negative")),
    ("Map rank", lambda b: (b.get("vis") or {}).get("rank") or ""),
    ("Profile score", lambda b: (b.get("gbp") or {}).get("score")),
    ("Profile claimed", lambda b: {True: "Yes", False: "No"}.get((b.get("gbp") or {}).get("claimed"), "Unknown")),
    ("Lead with", lambda b: SERVICE_LABELS.get((b.get("score") or {}).get("top"), "")),
    ("Headline", lambda b: (b.get("pitch") or {}).get("headline") or ""),
    ("Pitch", lambda b: (b.get("pitch") or {}).get("pitch") or ""),
    ("Email subject", lambda b: (b.get("pitch") or {}).get("email_subject") or ""),
    ("Call opener", lambda b: (b.get("pitch") or {}).get("call_opener") or ""),
    ("Evidence", lambda b: " | ".join(f["text"] for f in b.get("facts") or [])),
    ("Google Maps", lambda b: b["maps_url"]),
]


def _safe(v):
    """Neutralise spreadsheet formulas: a cell starting with = + - @ is text."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return "" if v is None else v


def lead_rows(payload):
    return [b for b in payload["businesses"] if b["researched"]]


def csv_text(payload):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c for c, _ in COLUMNS])
    for b in lead_rows(payload):
        w.writerow([_safe(f(b)) for _, f in COLUMNS])
    return buf.getvalue()


def xlsx_bytes(payload):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append([c for c, _ in COLUMNS])
    head = PatternFill("solid", fgColor="121213")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="F6F4F0")
        cell.fill = head
    tier_fill = {"A": "FFD2BF", "B": "FFE7A8", "C": "EEEAE4"}
    for b in lead_rows(payload):
        ws.append([_safe(f(b)) for _, f in COLUMNS])
        t = (b.get("score") or {}).get("tier")
        if t in tier_fill:
            ws.cell(row=ws.max_row, column=2).fill = PatternFill("solid", fgColor=tier_fill[t])
    widths = [6, 6, 10, 30, 18, 36, 16, 28, 30, 20, 8, 8, 10, 8, 8, 10, 26, 40, 70, 30, 50, 90, 40]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = wdt
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (18, 19, 22))
    ws.freeze_panes = "E2"
    found = wb.create_sheet("Everything found")
    found.append(["Business", "Category", "Address", "Phone", "Website", "Rating", "Reviews", "Researched",
                  "Chain", "Google Maps"])
    for b in payload["businesses"]:
        found.append([_safe(v) for v in (b["name"], b["category"], b["address"], b["phone"], b["website"],
                                         b.get("rating"), b.get("reviews"), "Yes" if b["researched"] else "No",
                                         b["chain_reason"] or ("" if not b["chain"] else "Yes"), b["maps_url"])])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
