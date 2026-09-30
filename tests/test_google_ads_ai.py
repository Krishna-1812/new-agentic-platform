"""Google Ads AI review (tracker/gads_ai.py, tracker/gads_ai_store.py and the routes in app.py).

Claude is replaced by a recording stand-in, so these check what the review reads, what it measures
itself, exactly what it asks Claude for, and how the page and its API behave; never a real call."""
import io
import json
import os
import types

import pytest

from tests import test_google_ads_insights as T
from tracker import gads_ai as ai
from tracker import gads_ai_store as store
from tracker import google_ads_insights as g

appmod = T.appmod

ACCOUNT = "Acme"


@pytest.fixture(autouse=True)
def _memory(monkeypatch):
    monkeypatch.delenv("DATABASE_URL", raising=False)
    store.reset_memory()
    yield
    store.reset_memory()


def _state():
    return g.load(T._raw(T._harness()))


REPORT_ROWS = [{"account": ACCOUNT, "customer_id": "123", "campaign": "Generic - Search", "state": "Enabled",
                "type": "Search", "day": "2026-09-26", "cost": 10.0, "currency": "INR", "clicks": 1.0,
                "impressions": 10.0, "conversions": 0.0, "cost_native": 10.0, "currency_native": "INR"}]


# ── A stand-in for the Anthropic client ─────────────────────────────────────
class FakeClaude:
    """Answers each streamed request with the next canned reply, and keeps what it was asked."""

    def __init__(self, *replies):
        self.replies = list(replies)
        self.calls = []
        self.beta = types.SimpleNamespace(messages=types.SimpleNamespace(stream=self._stream))

    def _stream(self, **kw):
        self.calls.append(kw)
        reply = self.replies.pop(0)
        return _Stream(reply)


class _Stream:
    def __init__(self, reply):
        self.reply = reply

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def get_final_message(self):
        r = self.reply
        if isinstance(r, Exception):
            raise r
        stop, body = (r if isinstance(r, tuple) else ("end_turn", r))
        text = body if isinstance(body, str) else json.dumps(body)
        return types.SimpleNamespace(
            stop_reason=stop, model=ai.MODEL,
            content=[types.SimpleNamespace(type="thinking", thinking=""), types.SimpleNamespace(type="text", text=text)],
            usage=types.SimpleNamespace(input_tokens=100000, output_tokens=20000, cache_read_input_tokens=0,
                                        cache_creation_input_tokens=0))


TARGETS = {"business": "CRM software", "objectives": ["Leads at a sustainable cost"],
           "targets": [{"metric": "cpa", "scope": "account", "comparison": "at_most", "value": 1000, "currency": "INR",
                        "items": [], "description": "Cost per lead at most INR 1,000", "quote": "CPA under 1000."},
                       {"metric": "roas", "scope": "account", "comparison": "at_least", "value": 400, "currency": "",
                        "items": [], "description": "ROAS 4x", "quote": "ROAS of 4x."},
                       {"metric": "locations_target", "scope": "account", "comparison": "only", "value": 0,
                        "currency": "", "items": ["Mumbai"], "description": "Only Mumbai", "quote": "Mumbai only."},
                       {"metric": "cpa", "scope": "Nonexistent campaign", "comparison": "at_most", "value": 1,
                        "currency": "", "items": [], "description": "x", "quote": "x"}],
           "account_manager_tasks": [{"task": "Review search terms", "cadence": "weekly", "quote": "Weekly."}],
           "brand_terms": ["acme"], "negative_themes": ["free"], "missing": []}

REVIEW = {"headline": "Leads cost more than the brief allows.", "overall": "needs_attention",
          "scorecard": [{"objective": "CPA", "target": "INR 1,000", "actual": "INR 1,333", "status": "off_track",
                         "evidence": "Generic - Search"}],
          "brief_compliance": [{"requirement": "Mumbai only", "status": "partly", "evidence": "x", "action": "y"}],
          "working": [{"title": "Brand", "evidence": "e", "where": "Brand - Search"}],
          "not_working": [{"title": "Waste", "evidence": "free crm software", "impact": "INR 10,800", "where": "Generic"}],
          "actions": [{"priority": "P1", "title": "Add negatives", "what_to_do": "Add 'free'", "why": "w",
                       "expected_impact": "i", "where": "Generic - Search", "effort": "low"}],
          "campaigns": [{"campaign": "Generic - Search", "verdict": "fix", "summary": "s", "issues": ["i"],
                         "opportunities": ["o"]}],
          "brief_gaps": ["No monthly budget."], "data_caveats": ["Rare terms are withheld."]}


# ── Briefs ───────────────────────────────────────────────────────────────────
def test_brief_files_are_read_as_text():
    assert ai.brief_text_from_file("b.md", "Target CPA ₹900".encode("utf-8")) == "Target CPA ₹900"
    import docx
    d = docx.Document()
    d.add_paragraph("Target ROAS 400%")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "City", "Pune"
    buf = io.BytesIO()
    d.save(buf)
    text = ai.brief_text_from_file("brief.docx", buf.getvalue())
    assert "Target ROAS 400%" in text and "City | Pune" in text, "tables come through as rows"
    from reportlab.pdfgen import canvas
    pdf = io.BytesIO()
    c = canvas.Canvas(pdf)
    c.drawString(72, 720, "Monthly budget 5 lakh")
    c.save()
    assert "Monthly budget 5 lakh" in ai.brief_text_from_file("b.pdf", pdf.getvalue())
    with pytest.raises(ai.ReviewError):
        ai.brief_text_from_file("b.xlsx", b"x")
    with pytest.raises(ai.ReviewError):
        ai.clean_brief("x" * (ai.BRIEF_MAX_CHARS + 1))


def test_briefs_are_versioned_and_reviews_remember_theirs():
    one = store.save_brief(ACCOUNT, "v1", "a@x.com", customer_id="123")
    two = store.save_brief(ACCOUNT, "v2", "b@x.com", filename="brief.docx")
    assert store.get_brief(ACCOUNT)["text"] == "v2"
    assert store.get_brief(brief_id=one["id"])["text"] == "v1"
    assert [b["id"] for b in store.brief_history(ACCOUNT)] == [two["id"], one["id"]]
    assert store.briefed_accounts().keys() == {ACCOUNT}


def test_a_review_whose_heartbeat_stopped_reads_as_failed():
    rid = store.create_review(ACCOUNT, "a@x.com")
    store.update_review(rid, status="running", heartbeat_at="2020-01-01T00:00:00+00:00")
    r = store.get_review(rid)
    assert r["status"] == "failed" and "restarted" in r["error"]
    assert store.active_review(ACCOUNT) is None


# ── The pack ─────────────────────────────────────────────────────────────────
def test_the_pack_has_every_campaign_with_figures_worked_out_in_code():
    pack = ai.build_pack(g, _state(), ACCOUNT, REPORT_ROWS, g.FX(REPORT_ROWS), g.Spend(REPORT_ROWS))
    assert pack["period"]["last_30_days"] == ["2026-08-28", "2026-09-26"], "30 days ending on the account's last day"
    names = {c["campaign"] for c in pack["campaigns"]}
    assert {"Brand - Search", "Generic - Search"} <= names
    gen = next(c for c in pack["campaigns"] if c["campaign"] == "Generic - Search")
    p = gen["last_30_days"]
    assert p["cpa"] == round(p["cost"] / p["conv"], 2) and p["ctr_pct"] == round(100 * p["clicks"] / p["impr"], 2)
    assert gen["bidding"] == "MAXIMIZE_CONVERSIONS" and gen["impression_share"]["lost_to_budget_pct"] == 90.0
    assert gen["search_terms"]["spend_without_conversions"][0]["term"] == "free crm software"
    assert gen["keywords"][0]["quality_score"] == 4 and gen["keywords_summary"]["quality_score_4_or_less"] >= 1
    assert gen["ads_summary"]["disapproved_or_limited"] >= 0
    assert pack["conversion_setup"] and pack["budgets_this_month"]
    assert pack["account_totals"]["last_30_days"]["cost"] == pytest.approx(sum(c["last_30_days"]["cost"] for c in pack["campaigns"]))
    json.dumps(pack)   # all of it is plain JSON


def test_an_account_without_daily_figures_is_refused_with_a_reason():
    with pytest.raises(ai.ReviewError) as e:
        ai.build_pack(g, g.load({}), ACCOUNT, REPORT_ROWS, g.FX(REPORT_ROWS))
    assert "insights script" in str(e.value)


# ── Checks made in code ──────────────────────────────────────────────────────
def test_targets_are_measured_in_code_before_claude_sees_them():
    pack = ai.build_pack(g, _state(), ACCOUNT, REPORT_ROWS, g.FX(REPORT_ROWS), g.Spend(REPORT_ROWS))
    checks = {(c["metric"], c["scope"]): c for c in ai.compute_checks(TARGETS, pack)}
    tot = pack["account_totals"]["last_30_days"]
    cpa = checks[("cpa", "account")]
    assert cpa["status"] == ai._judge(tot["cpa"], 1000, "at_most") and ai._fmt(tot["cpa"]) in cpa["actual"]
    roas = checks[("roas", "account")]
    assert roas["status"] == "cannot_verify" and "no conversion value" in roas["actual"], "no values: not measurable"
    loc = checks[("locations_target", "account")]
    assert "% of spend from people in the target places" in loc["actual"]
    assert checks[("cpa", "Nonexistent campaign")]["status"] == "cannot_verify"


def test_judging_a_number_against_a_target():
    assert ai._judge(900, 1000, "at_most") == "met"
    assert ai._judge(1050, 1000, "at_most") == "at_risk"
    assert ai._judge(1300, 1000, "at_most") == "off_track"
    assert ai._judge(380, 400, "at_least") == "at_risk"
    assert ai._judge(None, 400, "at_least") == "cannot_verify"


def test_a_money_target_in_the_accounts_own_currency_is_converted_at_googles_rate():
    rows = [dict(REPORT_ROWS[0], cost=830.0, cost_native=10.0, currency_native="USD")]
    fx = g.FX(rows)
    pack = {"currency": "INR", "campaigns": [{"campaign": "C", "last_30_days": ai._perf({"cost": 30000, "conv": 30}),
                                              "daily_budget": 1000, "shared_budget": False}],
            "budgets_this_month": [], "locations": {"costliest": [], "total": None}}
    t = {"targets": [{"metric": "cpa", "scope": "account", "comparison": "at_most", "value": 13, "currency": "USD",
                      "items": [], "description": "CPA $13", "quote": "q"}]}
    c = ai.compute_checks(t, pack, fx, ACCOUNT)[0]
    assert c["status"] == "met" and "converted at Google's latest rate, 83.0000" in c["actual"], "CPA 1000 against 13 x 83 = 1079"


# ── What Claude is asked ─────────────────────────────────────────────────────
def test_a_whole_review_runs_and_asks_claude_the_right_way():
    store.save_brief(ACCOUNT, "Brief: CPA under 1000. Mumbai only.", "a@x.com")
    rid = store.create_review(ACCOUNT, "a@x.com", brief_id=store.get_brief(ACCOUNT)["id"])
    fake = FakeClaude(TARGETS, REVIEW)
    st = _state()
    ai.run_review(rid, ACCOUNT, lambda: (g, st, REPORT_ROWS, g.FX(REPORT_ROWS), g.Spend(REPORT_ROWS), None),
                  store, client=fake, beat_every=0.05)
    r = store.get_review(rid)
    assert r["status"] == "complete", r["error"]
    assert r["report"]["headline"] == REVIEW["headline"] and r["report"]["has_brief"]
    assert r["report"]["checks"] and r["report"]["targets"]["targets"][0]["metric"] == "cpa"
    assert r["cost"]["usd"] == pytest.approx(2 * (0.1 * 4 + 0.02 * 20))
    assert r["period_from"] == "2026-08-28" and r["stats"]["campaigns"] >= 2
    for call, effort in zip(fake.calls, ("medium", "high")):
        assert call["model"] == "claude-opus-5-5"
        assert call["thinking"] == {"type": "adaptive"}
        assert call["output_config"]["effort"] == effort and call["output_config"]["format"]["type"] == "json_schema"
        assert call["betas"] == ["server-side-fallback-2026-07-01"] and call["fallbacks"] == "default"
        assert "temperature" not in call and "budget_tokens" not in json.dumps(call.get("thinking"))
    targets_call, review_call = fake.calls
    assert "<brief>\nBrief: CPA under 1000. Mumbai only.\n</brief>" in targets_call["messages"][0]["content"]
    body = review_call["messages"][0]["content"]
    assert "<measured_checks>" in body and "<data_pack>" in body and "Generic - Search" in body
    assert review_call["max_tokens"] == 64000


def test_structured_output_schemas_are_strict():
    def walk(s, path="$"):
        if s.get("type") == "object":
            assert s.get("additionalProperties") is False, path
            assert set(s["required"]) == set(s["properties"]), path
            for k, v in s["properties"].items():
                walk(v, path + "." + k)
        if s.get("type") == "array":
            walk(s["items"], path + "[]")
    walk(ai.REVIEW_SCHEMA)
    walk(ai.TARGETS_SCHEMA)


@pytest.mark.parametrize("reply, words", [
    (("refusal", ""), "declined"),
    (("max_tokens", "{"), "cut off"),
    (("end_turn", "not json"), "not readable"),
])
def test_a_review_that_cannot_finish_says_why(reply, words):
    rid = store.create_review(ACCOUNT, "a@x.com")
    st = _state()
    ai.run_review(rid, ACCOUNT, lambda: (g, st, REPORT_ROWS, g.FX(REPORT_ROWS), None, None), store,
                  client=FakeClaude(reply), beat_every=0.05)
    r = store.get_review(rid)
    assert r["status"] == "failed" and words in r["error"]


def test_no_brief_still_reviews_on_the_accounts_own_numbers():
    rid = store.create_review(ACCOUNT, "a@x.com")
    fake = FakeClaude(REVIEW)   # no brief: the extraction call is skipped
    st = _state()
    ai.run_review(rid, ACCOUNT, lambda: (g, st, REPORT_ROWS, g.FX(REPORT_ROWS), None, None), store,
                  client=fake, beat_every=0.05)
    r = store.get_review(rid)
    assert r["status"] == "complete" and not r["report"]["has_brief"] and len(fake.calls) == 1
    assert "There is no brief for this account yet" in fake.calls[0]["messages"][0]["content"]


def test_api_errors_become_plain_reasons():
    class E(Exception):
        status_code = 401
    rid = store.create_review(ACCOUNT, "a@x.com")
    st = _state()
    ai.run_review(rid, ACCOUNT, lambda: (g, st, REPORT_ROWS, g.FX(REPORT_ROWS), None, None), store,
                  client=FakeClaude(E("bad key")), beat_every=0.05)
    assert "refused the API key" in store.get_review(rid)["error"]


# ── The page and its API ─────────────────────────────────────────────────────
@pytest.fixture
def page(monkeypatch):
    st = _state()
    monkeypatch.setattr(appmod, "_fetch_google_ads_rows", lambda force=False: REPORT_ROWS)
    monkeypatch.setattr(appmod, "_gads_ai_load",
                        lambda: (g, st, REPORT_ROWS, g.FX(REPORT_ROWS), g.Spend(REPORT_ROWS), None))
    return T._client()


def test_the_page_lists_the_accounts_and_warns_what_is_missing(page, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    html = page.get("/dashboards/google-ads/ai-review").get_data(as_text=True)
    assert ">Acme (no brief)</option>" in html and "Claude is not configured" in html and "kept in memory only" in html
    assert 'id="gar-run" disabled' in html.replace("\n", " ").replace("  ", " ") or "disabled" in html
    dash = page.get("/dashboards/google-ads").get_data(as_text=True)
    assert 'id="gad-ai-link"' in dash


def test_saving_a_brief_needs_the_pages_own_request(page):
    r = page.post("/api/dashboards/google-ads/ai/brief", json={"account": ACCOUNT, "text": "x"})
    assert r.status_code == 400, "a plain cross-site POST is refused"
    h = {"X-Requested-With": "fetch"}
    assert page.post("/api/dashboards/google-ads/ai/brief", json={"account": "Other", "text": "x"}, headers=h).status_code == 400
    ok = page.post("/api/dashboards/google-ads/ai/brief", json={"account": ACCOUNT, "text": "CPA 900"}, headers=h).get_json()
    assert ok["ok"] and ok["brief"]["email"] == "reporting@markifydigital.com"
    up = page.post("/api/dashboards/google-ads/ai/brief", headers=h, content_type="multipart/form-data",
                   data={"account": ACCOUNT, "file": (io.BytesIO(b"ROAS 5x"), "brief.txt")}).get_json()
    assert up["ok"] and up["brief"]["filename"] == "brief.txt" and len(up["history"]) == 2
    got = page.get("/api/dashboards/google-ads/ai/brief?account=Acme").get_json()
    assert got["brief"]["text"] == "ROAS 5x"


def test_starting_a_review_runs_it_in_the_background(page, monkeypatch):
    h = {"X-Requested-With": "fetch"}
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    assert page.post("/api/dashboards/google-ads/ai/reviews", json={"account": ACCOUNT}, headers=h).status_code == 503
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
    fake = FakeClaude(REVIEW)
    real = ai.run_review
    monkeypatch.setattr(ai, "run_review", lambda *a, **k: real(*a, client=fake, beat_every=0.05, **k))
    started = []
    monkeypatch.setattr(appmod, "_gads_ai_spawn", lambda fn, *args: started.append(fn) or fn(*args))
    r = page.post("/api/dashboards/google-ads/ai/reviews", json={"account": ACCOUNT}, headers=h)
    assert r.status_code == 202 and len(started) == 1
    rid = r.get_json()["review"]["id"]
    got = page.get("/api/dashboards/google-ads/ai/reviews/%d" % rid).get_json()["review"]
    assert got["status"] == "complete" and got["report"]["overall"] == "needs_attention"
    listed = page.get("/api/dashboards/google-ads/ai/reviews?account=Acme").get_json()["reviews"]
    assert listed[0]["id"] == rid and "report" not in listed[0]


def test_one_review_at_a_time_per_account(page, monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
    rid = store.create_review(ACCOUNT, "a@x.com")
    r = page.post("/api/dashboards/google-ads/ai/reviews", json={"account": ACCOUNT},
                  headers={"X-Requested-With": "fetch"}).get_json()
    assert r["already"] and r["review"]["id"] == rid


def test_the_page_script_never_writes_server_text_as_html():
    with open(os.path.join(T.ROOT, "static", "js", "google-ads-ai.js"), encoding="utf-8") as fh:
        src = fh.read()
    assert "innerHTML" not in src and "insertAdjacentHTML" not in src
    assert '"X-Requested-With": "fetch"' in src
